from datetime import datetime

from openpyxl import load_workbook

from domain.models import DateRange, TrackPoint, Trip
from infrastructure.track_record_writer import TrackRecordWriter


def _trip() -> Trip:
    return Trip(
        ruc_empresa="20123456789",
        placa="A0P736",
        imei="863238075031916",
        codigo_ruta="R1",
        sentido=0,
        nro_doc_conductor="06583153",
        object_id="obj-1",
        date_range=DateRange(
            start=datetime(2026, 7, 6, 0, 0, 0),
            end=datetime(2026, 7, 7, 23, 59, 59),
        ),
    )


def _points() -> list[TrackPoint]:
    return [
        TrackPoint(
            datetime="2026-07-06T09:00:00",
            latitude=-12.0,
            longitude=-77.0,
            speed=10.0,
            ignition_status="ON",
        ),
        TrackPoint(
            datetime="2026-07-07T01:00:00",
            latitude=-12.1,
            longitude=-77.1,
            speed=20.0,
            ignition_status="ON",
        ),
    ]


def test_writer_emits_final_xlsx_format(tmp_path):

    out = tmp_path / "out.xlsx"

    writer = TrackRecordWriter(str(out), include_ignition_status=False)
    writer.open()
    writer.append(_trip(), _points())
    writer.close()

    wb = load_workbook(out)

    assert wb.sheetnames == ["2026-07-06", "2026-07-07"]

    sheet = wb["2026-07-06"]

    assert sheet.freeze_panes == "A2"

    header = [c.value for c in sheet[1]]
    row = [c.value for c in sheet[2]]

    assert header == TrackRecordWriter.BASE_HEADERS

    # Fechas sin la "T", como texto exacto.
    assert row[4] == "2026-07-06 00:00:00"
    assert row[5] == "2026-07-06 09:00:00"
    assert "T" not in row[5]

    # IDs como texto (sin perder ceros ni pasar a notacion cientifica).
    assert row[2] == "863238075031916" and isinstance(row[2], str)
    assert row[10] == "06583153" and isinstance(row[10], str)
    assert row[9] == "0" and isinstance(row[9], str)

    # Coordenadas y velocidad como numero (no texto).
    assert row[6] == -12.0 and isinstance(row[6], (int, float))
    assert row[7] == -77.0 and isinstance(row[7], (int, float))
    assert row[8] == 10.0 and isinstance(row[8], (int, float))


def test_writer_also_emits_formatted_csv(tmp_path):

    out = tmp_path / "out.xlsx"

    writer = TrackRecordWriter(str(out), include_ignition_status=False)
    writer.open()
    writer.append(_trip(), _points())
    writer.close()

    csv_path = tmp_path / "out.csv"

    assert writer.csv_file == csv_path
    assert csv_path.exists()

    lines = csv_path.read_text(encoding="utf-8").splitlines()

    # Cabecera + las dos filas de ambos dias, todo en un unico CSV.
    assert lines[0] == "|".join(TrackRecordWriter.BASE_HEADERS)
    assert len(lines) == 3

    first = lines[1].split("|")

    assert first[4] == "2026-07-06 00:00:00"
    assert first[5] == "2026-07-06 09:00:00"
    assert first[2] == "863238075031916"
    assert first[10] == "06583153"

    assert lines[2].split("|")[5] == "2026-07-07 01:00:00"


def test_writer_can_skip_csv(tmp_path):

    out = tmp_path / "out.xlsx"

    writer = TrackRecordWriter(str(out), write_csv=False)
    writer.open()
    writer.append(_trip(), _points())
    writer.close()

    assert writer.csv_file is None
    assert not (tmp_path / "out.csv").exists()


def test_writer_keeps_ignition_column_when_enabled(tmp_path):

    out = tmp_path / "out.xlsx"

    writer = TrackRecordWriter(str(out), include_ignition_status=True)
    writer.open()
    writer.append(_trip(), _points())
    writer.close()

    wb = load_workbook(out)
    sheet = wb["2026-07-06"]

    header = [c.value for c in sheet[1]]
    row = [c.value for c in sheet[2]]

    assert header[-1] == "IGNITION_STATUS"
    assert row[-1] == "ON"

    csv_lines = (tmp_path / "out.csv").read_text(encoding="utf-8").splitlines()
    assert csv_lines[0].endswith("|IGNITION_STATUS")
    assert csv_lines[1].split("|")[-1] == "ON"
