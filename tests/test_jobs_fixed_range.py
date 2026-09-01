import pandas as pd
from openpyxl import load_workbook

from application.jobs.fixed_range_job import FixedRangeJob
from domain.models import Vehicle
from tests.conftest import FakeFMTrackClient


def test_fixed_range_job_applies_same_range_to_every_row(tmp_path):

    vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="obj-1",
            imei="863238075031916",
            credential="default",
        )
    }

    pages = {
        "obj-1": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                        "ignition_status": "ON",
                    },
                ],
                "continuation_token": None,
            },
        ]
    }

    client = FakeFMTrackClient(objects=vehicles, coordinate_pages=pages)

    listing = pd.DataFrame(
        {
            "RUC": ["123"],
            "CODIGO_RUTA": ["R1"],
            "PLACA": ["A0P736"],
            "IMEI": ["863238075031916"],
        }
    )

    input_file = tmp_path / "listing.xlsx"
    output_file = tmp_path / "out.xlsx"

    listing.to_excel(input_file, index=False)

    job = FixedRangeJob(client)

    summary = job.run(
        str(output_file),
        input_file=str(input_file),
        start_date="2026-07-06 00:00:00",
        end_date="2026-07-06 23:59:59",
    )

    assert summary.processed == 1
    assert summary.track_points == 1

    workbook = load_workbook(output_file)

    assert "2026-07-06" in workbook.sheetnames


class FakeVehicleSourceClient:
    """
    Doble de prueba de LucaClient para el rol de vehicle_source_client. No hace red.
    """

    def __init__(self, objects: dict[str, Vehicle]):
        self._objects = objects

    def get_objects(self) -> dict[str, Vehicle]:
        return dict(self._objects)


def test_fixed_range_job_builds_listing_from_luca_when_no_input_file(tmp_path):

    luca_vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="A0P736",
            imei="863238075031916",
            credential="luca",
        )
    }

    fmtrack_vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="obj-1",
            imei="863238075031916",
            credential="default",
        )
    }

    pages = {
        "obj-1": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                        "ignition_status": "ON",
                    },
                ],
                "continuation_token": None,
            },
        ]
    }

    vehicle_source_client = FakeVehicleSourceClient(luca_vehicles)
    track_client = FakeFMTrackClient(objects=fmtrack_vehicles, coordinate_pages=pages)

    output_file = tmp_path / "out.xlsx"

    job = FixedRangeJob(track_client, vehicle_source_client=vehicle_source_client)

    summary = job.run(
        str(output_file),
        start_date="2026-07-06 00:00:00",
        end_date="2026-07-06 23:59:59",
        codigo_ruta=6,
    )

    assert summary.processed == 1
    assert summary.track_points == 1

    workbook = load_workbook(output_file)

    assert "2026-07-06" in workbook.sheetnames
