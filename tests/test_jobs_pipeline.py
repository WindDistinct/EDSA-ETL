import pandas as pd
from openpyxl import load_workbook

from application.jobs.pipeline_job import PipelineJob
from domain.models import Vehicle
from tests.conftest import FakeFMTrackClient


def test_pipeline_job_resolves_date_range_per_row(tmp_path):

    vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="obj-1",
            imei="863238075031916",
            credential="default",
        )
    }

    pages = {
        "obj-1": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                        "ignition_status": "ON",
                    },
                ],
                "continuation_token": None,
            },
        ]
    }

    client = FakeFMTrackClient(objects=vehicles, coordinate_pages=pages)

    df = pd.DataFrame(
        {
            "RUC_EMPRESA": ["123"],
            "PLACA": ["A0P736"],
            "IMEI": ["863238075031916"],
            "CODIGO_RUTA": ["R1"],
            "SENTIDO": ["0"],
            "NRO_DOC_CONDUCTOR": ["12345678"],
            "FECHORA_INI_VIAJE": ["2026-07-06T08:00:00"],
            "FECHA_HORA_FIN_TRAMO": ["2026-07-06T20:00:00"],
        }
    )

    input_file = tmp_path / "partition.xlsx"
    output_file = tmp_path / "out.xlsx"

    df.to_excel(input_file, index=False)

    job = PipelineJob(client)

    summary = job.run(str(input_file), str(output_file))

    assert summary.processed == 1
    assert summary.track_points == 1

    workbook = load_workbook(output_file)

    assert "2026-07-06" in workbook.sheetnames
