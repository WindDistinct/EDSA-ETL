from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from application.track_pipeline import TrackPipeline
from domain.models import DateRange, Vehicle
from tests.conftest import FakeFMTrackClient


def test_track_pipeline_writes_day_sheets_and_summary(tmp_path):

    vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="obj-1",
            imei="863238075031916",
            credential="default",
        )
    }

    # Lima es UTC-5: 14:00Z 07-06 -> 09:00 local 07-06; 06:00Z 07-07 -> 01:00 local 07-07
    pages = {
        "obj-1": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                        "ignition_status": "ON",
                    },
                    {
                        "datetime": "2026-07-07T06:00:00.000Z",
                        "position": {"latitude": -12.1, "longitude": -77.1, "speed": 20.0},
                        "ignition_status": "ON",
                    },
                ],
                "continuation_token": None,
            },
        ]
    }

    client = FakeFMTrackClient(objects=vehicles, coordinate_pages=pages)

    df = pd.DataFrame(
        {
            "PLACA": ["A0P736"],
            "RUC_EMPRESA": ["123"],
            "IMEI": ["863238075031916"],
            "CODIGO_RUTA": ["R1"],
            "SENTIDO": ["0"],
            "NRO_DOC_CONDUCTOR": ["12345678"],
        }
    )

    fixed_range = DateRange(
        start=datetime(2026, 7, 6, 0, 0, 0),
        end=datetime(2026, 7, 7, 23, 59, 59),
    )

    output_file = tmp_path / "out.xlsx"

    pipeline = TrackPipeline(client, include_ignition_status=True)

    summary = pipeline.run(
        df,
        str(output_file),
        date_range_fn=lambda row: fixed_range,
    )

    assert summary.total == 1
    assert summary.processed == 1
    assert summary.errors == 0
    assert summary.track_points == 2

    workbook = load_workbook(output_file)

    assert workbook.sheetnames == ["2026-07-06", "2026-07-07"]

    sheet_1 = workbook["2026-07-06"]

    header = [cell.value for cell in sheet_1[1]]
    row = [cell.value for cell in sheet_1[2]]

    assert header[1] == "PLACA"
    assert header[-1] == "IGNITION_STATUS"

    assert row[1] == "A0P736"
    assert row[6] == -12.0
    assert row[-1] == "ON"


def test_track_pipeline_counts_missing_object_id_as_error(tmp_path):

    client = FakeFMTrackClient(objects={})

    df = pd.DataFrame(
        {
            "PLACA": ["ZZZ999"],
            "RUC_EMPRESA": ["123"],
            "IMEI": ["000"],
            "CODIGO_RUTA": ["R1"],
            "SENTIDO": ["0"],
            "NRO_DOC_CONDUCTOR": ["12345678"],
        }
    )

    fixed_range = DateRange(
        start=datetime(2026, 7, 6, 0, 0, 0),
        end=datetime(2026, 7, 6, 23, 59, 59),
    )

    pipeline = TrackPipeline(client)

    summary = pipeline.run(
        df,
        str(tmp_path / "out.xlsx"),
        date_range_fn=lambda row: fixed_range,
    )

    assert summary.processed == 0
    assert summary.errors == 1
    assert summary.track_points == 0


def test_track_pipeline_resolves_trips_concurrently(tmp_path):

    plates = [f"P{i:03d}" for i in range(20)]

    vehicles = {
        plate: Vehicle(plate=plate, object_id=f"obj-{plate}", imei="123", credential="default")
        for plate in plates
    }

    pages = {
        f"obj-{plate}": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                        "ignition_status": "ON",
                    },
                ],
                "continuation_token": None,
            },
        ]
        for plate in plates
    }

    client = FakeFMTrackClient(objects=vehicles, coordinate_pages=pages)

    df = pd.DataFrame(
        {
            "PLACA": plates,
            "RUC_EMPRESA": ["123"] * len(plates),
            "IMEI": ["123"] * len(plates),
            "CODIGO_RUTA": ["R1"] * len(plates),
            "SENTIDO": ["0"] * len(plates),
            "NRO_DOC_CONDUCTOR": ["12345678"] * len(plates),
        }
    )

    fixed_range = DateRange(
        start=datetime(2026, 7, 6, 0, 0, 0),
        end=datetime(2026, 7, 6, 23, 59, 59),
    )

    pipeline = TrackPipeline(client, include_ignition_status=True, max_workers=8)

    summary = pipeline.run(
        df,
        str(tmp_path / "out.xlsx"),
        date_range_fn=lambda row: fixed_range,
    )

    assert summary.total == len(plates)
    assert summary.processed == len(plates)
    assert summary.errors == 0
    assert summary.track_points == len(plates)
