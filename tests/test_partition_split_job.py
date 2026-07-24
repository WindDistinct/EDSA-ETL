import pandas as pd
from openpyxl import load_workbook

from application.jobs.partition_split_job import PartitionSplitJob


def test_partition_split_creates_expected_number_of_files_with_daily_sheets(tmp_path):

    df = pd.DataFrame(
        {
            "RUC_EMPRESA": ["123"] * 4,
            "PLACA": ["A0P736"] * 4,
            "IMEI": ["863238075031916"] * 4,
            "CODIGO_RUTA": ["R1"] * 4,
            "NRO_DOC_CONDUCTOR": ["12345678"] * 4,
            "FECHORA_INI_VIAJE": [
                "2026-01-01 08:00:00",
                "2026-01-05 08:00:00",
                "2026-01-10 08:00:00",
                "2026-01-15 08:00:00",
            ],
        }
    )

    input_file = tmp_path / "input.xlsx"
    df.to_excel(input_file, index=False)

    output_folder = tmp_path / "partitions"

    PartitionSplitJob(partitions=2).run(
        str(input_file),
        str(output_folder),
    )

    files = sorted(output_folder.glob("*.xlsx"))

    assert len(files) == 2

    workbook = load_workbook(files[0])

    assert len(workbook.sheetnames) >= 1
    assert workbook.sheetnames[0] in {"2026-01-01", "2026-01-05"}
