from __future__ import annotations

from openpyxl import load_workbook

from application.jobs.dispatch_job import DispatchJob
from domain.models import Vehicle
from tests.conftest import FakeFMTrackClient


class FakeDispatchClient:
    """
    Doble de prueba de LucaClient para el rol de dispatch_client. No hace red.
    """

    def __init__(self, despachos: list[dict]):
        self._despachos = despachos
        self.calls: list[tuple[str, str]] = []

    def get_dispatch_records(self, fecha_inicio: str, fecha_fin: str) -> list[dict]:
        self.calls.append((fecha_inicio, fecha_fin))
        return list(self._despachos)


DESPACHO = {
    "ruc_empresa": "20123456789",
    "placa": "A0P736",
    "codigo_ruta": "R1",
    "fecha": "2026-07-06",
    "hora": "08:00:00",
    "hora_real_ultimo_punto": "Mon Jul 06 2026 14:00:00 GMT-0500 (hora estándar de Perú)",
    "hora_calculada_ultimo_punto": None,
    "sentido": "ida",
    "conductor": "PEREZ, JUAN, 12345678",
}


def test_dispatch_job_uses_dispatch_client_for_despachos_and_track_client_for_gps(tmp_path):

    vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="obj-1",
            imei="863238075031916",
            credential="default",
        )
    }

    pages = {
        "obj-1": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                    },
                ],
                "continuation_token": None,
            },
        ]
    }

    dispatch_client = FakeDispatchClient([DESPACHO])
    track_client = FakeFMTrackClient(objects=vehicles, coordinate_pages=pages)

    output_file = tmp_path / "out.xlsx"

    summary = DispatchJob(dispatch_client, track_client).run(
        fecha_inicio="2026-07-01",
        fecha_fin="2026-07-10",
        output_file=str(output_file),
    )

    assert dispatch_client.calls == [("2026-07-01", "2026-07-10")]
    assert summary.processed == 1
    assert summary.track_points == 1

    workbook = load_workbook(output_file)

    assert "2026-07-06" in workbook.sheetnames


def test_dispatch_job_works_when_dispatch_and_track_client_are_the_same_object(tmp_path):

    class FakeLucaLikeClient(FakeDispatchClient, FakeFMTrackClient):
        def __init__(self, despachos, objects, coordinate_pages):
            FakeDispatchClient.__init__(self, despachos)
            FakeFMTrackClient.__init__(self, objects=objects, coordinate_pages=coordinate_pages)

    vehicles = {
        "A0P736": Vehicle(
            plate="A0P736",
            object_id="A0P736",
            imei="863238075031916",
            credential="luca",
        )
    }

    pages = {
        "A0P736": [
            {
                "items": [
                    {
                        "datetime": "2026-07-06T14:00:00.000Z",
                        "position": {"latitude": -12.0, "longitude": -77.0, "speed": 10.0},
                    },
                ],
                "continuation_token": None,
            },
        ]
    }

    client = FakeLucaLikeClient([DESPACHO], vehicles, pages)

    output_file = tmp_path / "out.xlsx"

    summary = DispatchJob(client, client).run(
        fecha_inicio="2026-07-01",
        fecha_fin="2026-07-10",
        output_file=str(output_file),
    )

    assert summary.processed == 1
    assert summary.track_points == 1
