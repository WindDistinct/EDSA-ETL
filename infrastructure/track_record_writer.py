from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from domain.models import TrackPoint, Trip

DATE_FMT = "%Y-%m-%d %H:%M:%S"


def _as_text(value) -> str:
    """
    Fuerza el valor a texto para no perder ceros a la izquierda
    (NRO_DOC_CONDUCTOR = "06583153") ni pasar el IMEI a notacion cientifica.
    """

    if value is None:
        return ""

    return str(value)


def _as_datetime_text(value) -> str:
    """
    Normaliza una fecha a "YYYY-MM-DD HH:MM:SS" (sin la "T" de ISO 8601).
    """

    if value is None or value == "":
        return ""

    if isinstance(value, datetime):
        return value.strftime(DATE_FMT)

    text = str(value).strip()

    try:
        return datetime.fromisoformat(
            text.replace("Z", "+00:00")
        ).strftime(DATE_FMT)
    except ValueError:
        return text.replace("T", " ")


class TrackRecordWriter:

    BASE_HEADERS = [
        "RUC_EMPRESA",
        "PLACA",
        "IMEI",
        "CODIGO_RUTA",
        "FECHORA_INI_VIAJE",
        "FECHA_HORA_TRACK",
        "LATITUD",
        "LONGITUD",
        "VELOCIDAD",
        "SENTIDO",
        "NRO_DOC_CONDUCTOR",
    ]

    def __init__(
        self,
        output_file: str,
        include_ignition_status: bool = False,
        *,
        write_csv: bool = True,
        csv_file: str | None = None,
        csv_sep: str = "|",
    ):

        self.output_file = Path(output_file)

        if csv_file is not None:
            self.csv_file: Path | None = Path(csv_file)
        elif write_csv:
            self.csv_file = self.output_file.with_suffix(".csv")
        else:
            self.csv_file = None

        self.csv_sep = csv_sep

        self.workbook: Workbook | None = None

        self.sheets: dict[str, any] = {}

        self._csv_handle = None

        self._csv_writer = None

        self.include_ignition_status = include_ignition_status

        self.headers = self.BASE_HEADERS + (
            ["IGNITION_STATUS"]
            if include_ignition_status
            else []
        )

    def open(self):

        if self.workbook is not None:
            return

        self.output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        self.workbook = Workbook(write_only=True)

        if self.csv_file is not None:

            self.csv_file.parent.mkdir(parents=True, exist_ok=True)

            self._csv_handle = open(
                self.csv_file,
                "w",
                newline="",
                encoding="utf-8",
            )

            self._csv_writer = csv.writer(
                self._csv_handle,
                delimiter=self.csv_sep,
                lineterminator="\n",
            )

            self._csv_writer.writerow(self.headers)

    def _get_sheet(
        self,
        day: str,
    ):

        if day in self.sheets:
            return self.sheets[day]

        ws = self.workbook.create_sheet(day)

        ws.freeze_panes = "A2"

        ws.append(self.headers)

        self.sheets[day] = ws

        return ws

    def close(self):

        if self.workbook is not None:

            if not self.sheets:
                self.workbook.create_sheet("SIN_DATOS")

            self.workbook.save(
                self.output_file
            )

            self.workbook.close()

            self.workbook = None

        if self._csv_handle is not None:

            self._csv_handle.close()

            self._csv_handle = None

            self._csv_writer = None

    def append(
        self,
        trip: Trip,
        points: list[TrackPoint],
    ) -> None:

        fechora_ini = _as_datetime_text(trip.date_range.start)

        for point in points:

            day = point.datetime[:10]

            sheet = self._get_sheet(day)

            row = [
                _as_text(trip.ruc_empresa),
                _as_text(trip.placa),
                _as_text(trip.imei),
                _as_text(trip.codigo_ruta),
                fechora_ini,
                _as_datetime_text(point.datetime),
                point.latitude,
                point.longitude,
                point.speed,
                _as_text(trip.sentido),
                _as_text(trip.nro_doc_conductor),
            ]

            if self.include_ignition_status:
                row.append(
                    _as_text(point.ignition_status)
                )

            sheet.append(row)

            if self._csv_writer is not None:
                self._csv_writer.writerow(row)
